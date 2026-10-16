"""
รายงานครุภัณฑ์ตัดจำหน่าย รายละเอียด ตามแหล่งเงิน (รวมทุกรายการที่ตัดจำหน่ายแล้ว)
เลย์เอาต์เดียวกับรายงานค่าเสื่อมราคารายละเอียดตามแหล่งเงิน (depreciation_builder.build_detail_excel_bytes)
แต่ตัดคอลัมน์ค่าเสื่อม/มูลค่าสุทธิออก แทนด้วยประเภทครุภัณฑ์ สถานะ (เหตุผลตัดจำหน่าย) และวันที่ตัดจำหน่าย
"""
import io
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from db_connection import run_query
from depreciation_builder import (
    THAI_FONT,
    THAI_MONTHS,
    _BORDER,
    _GROUP_FILL,
    _HEADER_FILL,
    _SUBGROUP_FILL,
    fiscal_be_label,
    fiscal_year_and_column,
    to_thai_date,
)


def fetch_disposal_data(division=None, dept=None, section=None) -> pd.DataFrame:
    conditions = [
        "cm.DISPOSDATETIME IS NOT NULL",
        "cm.DISPOSCODE IS NOT NULL",
        # if both fields are set, the asset is disposed/transferred-out, full stop - don't
        # look at ASMSTTF transfer history at all. A later ASMSTTF row isn't reliable
        # evidence of physical reuse here (confirmed on real disposed/transferred-out
        # assets that still had later transfer rows with no bearing on whether the asset
        # came back into use)
    ]
    params = []

    if division:
        conditions.append("cm.LOCATEDIVISION = ?")
        params.append(division)
    if dept:
        conditions.append("cm.LOCATEDEPT = ?")
        params.append(dept)
    if section:
        conditions.append("cm.LOCATESECTION = ?")
        params.append(section)

    where_clause = " AND ".join(conditions)
    sql = f"""
    SELECT
        cm.ASSETCODE,
        cm.SUFFIX,
        dbo.GetSSBName(ISNULL(cm.THAINAME, m.THAINAME)) AS AssetName,
        COALESCE(cm.ACQDATETIME, m.FIRSTDATETIME) AS AcqDate,
        cm.QTY,
        cm.PRICE,
        cm.DISPOSDATETIME,
        cm.DISPOSCODE,
        dbo.GetSSBName(dc.THAINAME) AS DisposeReasonName,
        cm.LOCATEDIVISION,
        div.DivisionName,
        cm.ARTICLEGROUP,
        dbo.GetSSBName(ag.THAINAME) AS ArticleGroupName,
        cm.PURCHASEBUDGETCODE,
        dbo.GetSSBName(bg.THAINAME) AS BudgetName,
        dep.DEPREPERCENT,
        dep.SCRAPVALUE,
        dep.TOTALVALUEDEPRE AS AccumDepreAtDispose
    FROM ASMSTCM cm
    JOIN ASMST m ON m.ASSETCODE = cm.ASSETCODE
    LEFT JOIN Division div ON cm.LOCATEDIVISION = div.Division
    LEFT JOIN SYSCONFIG ag ON ag.CODE = cm.ARTICLEGROUP AND ag.CTRLCODE = 100012
    LEFT JOIN SYSCONFIG bg ON bg.CODE = cm.PURCHASEBUDGETCODE AND bg.CTRLCODE = 120010
    LEFT JOIN SYSCONFIG dc ON dc.CODE = cm.DISPOSCODE AND dc.CTRLCODE = 100011
    -- DEPREGROUP='1' = ขึ้นบัญชีสินทรัพย์ (คิดค่าเสื่อม) ส่วนครุภัณฑ์ต่ำกว่าเกณฑ์ไม่มีแถวนี้
    -- (ตัดเป็นค่าใช้จ่ายทันทีตอนซื้อ ถือว่าค่าเสื่อมสะสม = มูลค่าเต็ม)
    LEFT JOIN ASMSTDEP dep ON dep.ASSETCODE = cm.ASSETCODE AND dep.SUFFIX = cm.SUFFIX AND dep.DEPREGROUP = '1'
    WHERE {where_clause}
    ORDER BY cm.PURCHASEBUDGETCODE, cm.ARTICLEGROUP, cm.DISPOSDATETIME
    """
    df = run_query(sql, params=tuple(params))
    if df.empty:
        return df

    df["TotalValue"] = df["QTY"].fillna(0.0) * df["PRICE"].fillna(0.0)
    has_depre_row = df["AccumDepreAtDispose"].notna()
    df["AccumDepre"] = df["AccumDepreAtDispose"].fillna(0.0)
    df.loc[~has_depre_row, "AccumDepre"] = df.loc[~has_depre_row, "TotalValue"]
    df["NetValue"] = df["TotalValue"] - df["AccumDepre"]
    df["UsefulLifeYears"] = (100.0 / df["DEPREPERCENT"].replace(0, pd.NA)).round()
    df["AcqDateThai"] = df["AcqDate"].apply(to_thai_date)
    df["DisposeDateThai"] = df["DISPOSDATETIME"].apply(to_thai_date)
    df["DisposeReasonLabel"] = df["DisposeReasonName"].fillna(df["DISPOSCODE"])
    df["BudgetLabel"] = df["BudgetName"].fillna("(ไม่ระบุแหล่งเงิน)")
    df["ArticleGroupLabel"] = (
        "(" + df["ARTICLEGROUP"].fillna("") + ") " + df["ArticleGroupName"].fillna("")
    )
    df["DivisionLabel"] = df["DivisionName"].fillna("(ไม่ระบุฝ่าย)")
    return df


_COLUMNS = ["รหัส", "รายการ", "วันเดือนปี", "จำนวน", "มูลค่ารวม", "ประเภทครุภัณฑ์", "สถานะ", "วันที่อัพเดทสถานะ"]
_WIDTHS = [16, 34, 14, 8, 14, 26, 20, 16]


def build_excel_bytes(df: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "ตัดจำหน่าย"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    n_cols = len(_COLUMNS)
    for i, w in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    c = ws.cell(1, 1, "รายงานครุภัณฑ์ตัดจำหน่าย รายละเอียด ตามแหล่งเงิน")
    c.font = Font(name=THAI_FONT, size=18, bold=True)
    c.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    c2 = ws.cell(2, 1, "รายการที่ตัดจำหน่ายแล้วทั้งหมด")
    c2.font = Font(name=THAI_FONT, size=13)
    c2.alignment = Alignment(horizontal="center")
    row = 4

    if df.empty:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
        ws.cell(row, 1, "ไม่พบรายการตัดจำหน่ายตามเงื่อนไขที่เลือก")
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    grand_total_value = 0.0
    grand_count = 0

    for budget_label, bdf in df.groupby("BudgetLabel", sort=False):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
        gc = ws.cell(row, 1, budget_label)
        gc.font = Font(name=THAI_FONT, size=14, bold=True)
        gc.fill = _GROUP_FILL
        row += 1

        for art_label, adf in bdf.groupby("ArticleGroupLabel", sort=False):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
            ac = ws.cell(row, 1, art_label)
            ac.font = Font(name=THAI_FONT, size=13, bold=True)
            ac.fill = _SUBGROUP_FILL
            row += 1

            for col_idx, label in enumerate(_COLUMNS, start=1):
                cell = ws.cell(row, col_idx, label)
                cell.font = Font(name=THAI_FONT, size=12, bold=True)
                cell.fill = _HEADER_FILL
                cell.alignment = Alignment(horizontal="center")
                cell.border = _BORDER
            row += 1

            sub_total_value = 0.0
            for _, r in adf.iterrows():
                values = [
                    r["ASSETCODE"], r["AssetName"], r["AcqDateThai"], r["QTY"], r["TotalValue"],
                    r["ArticleGroupLabel"], r["DisposeReasonLabel"], r["DisposeDateThai"],
                ]
                for col_idx, value in enumerate(values, start=1):
                    cell = ws.cell(row, col_idx, value)
                    cell.font = Font(name=THAI_FONT, size=12)
                    cell.border = _BORDER
                    if col_idx == 5:
                        cell.number_format = "#,##0.00"
                        cell.alignment = Alignment(horizontal="right")
                    elif col_idx in (4, 8):
                        cell.alignment = Alignment(horizontal="center")
                sub_total_value += float(r["TotalValue"])
                row += 1

            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            tc = ws.cell(row, 1, f"รวม {art_label} ({len(adf)} รายการ)")
            tc.font = Font(name=THAI_FONT, size=12, bold=True)
            tc.alignment = Alignment(horizontal="right")
            c = ws.cell(row, 5, sub_total_value)
            c.font = Font(name=THAI_FONT, size=12, bold=True)
            c.number_format = "#,##0.00"
            c.alignment = Alignment(horizontal="right")
            row += 2

            grand_total_value += sub_total_value
            grand_count += len(adf)

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    gcell = ws.cell(row, 1, f"รวมทั้งสิ้น ({grand_count} รายการ)")
    gcell.font = Font(name=THAI_FONT, size=14, bold=True)
    gcell.alignment = Alignment(horizontal="right")
    gc2 = ws.cell(row, 5, grand_total_value)
    gc2.font = Font(name=THAI_FONT, size=14, bold=True)
    gc2.number_format = "#,##0.00"
    gc2.alignment = Alignment(horizontal="right")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


_SUMMARY_COLUMNS = ["ลำดับ", "รายการ", "จำนวน", "มูลค่ารวม", "อายุการใช้งาน (ปี)", "ค่าเสื่อม สะสม", "มูลค่าสุทธิ"]
_SUMMARY_WIDTHS = [8, 34, 10, 16, 14, 16, 16]


def build_source_summary_excel_bytes(df: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "ตัดจำหน่าย สรุป"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    n_cols = len(_SUMMARY_COLUMNS)
    for i, w in enumerate(_SUMMARY_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    c = ws.cell(1, 1, "รายงานครุภัณฑ์ตัดจำหน่าย สรุปตามแหล่งเงิน/ประเภท")
    c.font = Font(name=THAI_FONT, size=18, bold=True)
    c.alignment = Alignment(horizontal="center")

    today = date.today()
    asmst_year, _ = fiscal_year_and_column(today.year, today.month)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    c2 = ws.cell(2, 1, f"รายงานทรัพย์สิน ณ สิ้นเดือน{THAI_MONTHS[today.month]} {fiscal_be_label(asmst_year)}")
    c2.font = Font(name=THAI_FONT, size=13)
    c2.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=n_cols)
    c3 = ws.cell(3, 1, "รายการที่ตัดจำหน่ายแล้วทั้งหมด")
    c3.font = Font(name=THAI_FONT, size=13)
    c3.alignment = Alignment(horizontal="center")
    row = 5

    if df.empty:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
        ws.cell(row, 1, "ไม่พบรายการตัดจำหน่ายตามเงื่อนไขที่เลือก")
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    grand_count = 0
    grand_value = 0.0
    grand_accum = 0.0
    grand_net = 0.0

    for budget_label, bdf in df.groupby("BudgetLabel", sort=False):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
        gc = ws.cell(row, 1, budget_label)
        gc.font = Font(name=THAI_FONT, size=14, bold=True)
        gc.fill = _GROUP_FILL
        row += 1

        for col_idx, label in enumerate(_SUMMARY_COLUMNS, start=1):
            cell = ws.cell(row, col_idx, label)
            cell.font = Font(name=THAI_FONT, size=12, bold=True)
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = _BORDER
        row += 1

        agg = bdf.groupby("ArticleGroupLabel", sort=False).agg(
            count=("ASSETCODE", "count"),
            TotalValue=("TotalValue", "sum"),
            UsefulLifeYears=("UsefulLifeYears", "first"),
            AccumDepre=("AccumDepre", "sum"),
            NetValue=("NetValue", "sum"),
        ).reset_index()

        for seq, (_, r) in enumerate(agg.iterrows(), start=1):
            values = [
                seq, r["ArticleGroupLabel"], r["count"], r["TotalValue"],
                r["UsefulLifeYears"], r["AccumDepre"], r["NetValue"],
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row, col_idx, value)
                cell.font = Font(name=THAI_FONT, size=12)
                cell.border = _BORDER
                if col_idx in (1, 3):
                    cell.alignment = Alignment(horizontal="center")
                elif col_idx >= 4:
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
            row += 1

        sub_count = int(agg["count"].sum())
        sub_value = float(agg["TotalValue"].sum())
        sub_accum = float(agg["AccumDepre"].sum())
        sub_net = float(agg["NetValue"].sum())

        tc = ws.cell(row, 1, f"รวม {budget_label}")
        tc.font = Font(name=THAI_FONT, size=12, bold=True)
        tc.fill = _GROUP_FILL
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        for col_idx, val in [(3, sub_count), (4, sub_value), (6, sub_accum), (7, sub_net)]:
            c = ws.cell(row, col_idx, val)
            c.font = Font(name=THAI_FONT, size=12, bold=True)
            c.fill = _GROUP_FILL
            if col_idx == 3:
                c.alignment = Alignment(horizontal="center")
            else:
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")
        ws.cell(row, 5).fill = _GROUP_FILL
        row += 1

        grand_count += sub_count
        grand_value += sub_value
        grand_accum += sub_accum
        grand_net += sub_net

    gcell = ws.cell(row, 1, "รวมทั้งสิ้น")
    gcell.font = Font(name=THAI_FONT, size=14, bold=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    for col_idx, val in [(3, grand_count), (4, grand_value), (6, grand_accum), (7, grand_net)]:
        c = ws.cell(row, col_idx, val)
        c.font = Font(name=THAI_FONT, size=14, bold=True)
        if col_idx == 3:
            c.alignment = Alignment(horizontal="center")
        else:
            c.number_format = "#,##0.00"
            c.alignment = Alignment(horizontal="right")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
