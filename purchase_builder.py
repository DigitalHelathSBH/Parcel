"""
รายงานซื้อประจำเดือน - ใช้เตรียมข้อมูลก่อนกรอกแบบ สขร.1
ดึงจาก SKPO/SKPODTL (ใบสั่งซื้อจริงของระบบพัสดุ) ตามช่วงวันที่ออกใบสั่งซื้อ (ISSUEDATETIME)
ในเดือนที่เลือก แต่ละใบสั่งซื้อแสดงเป็นหัวข้อ ตามด้วยตารางรายการย่อยแยกรายบรรทัด

หมายเหตุ: "วิธีซื้อหรือจ้าง" มาจาก SKPO.PURCHASETYPECODE (SYSCONFIG CTRLCODE=100010)
ซึ่งเป็นข้อมูลจริงที่ระบบเก็บไว้ (ไม่ใช่ช่องว่างให้กรอกเองเหมือนก่อนหน้านี้)
"""
import calendar
import io
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from db_connection import run_query
from depreciation_builder import THAI_FONT, THAI_MONTHS, _BORDER, _GROUP_FILL, _HEADER_FILL, to_thai_date


def fetch_purchase_data(calendar_year: int, calendar_month: int) -> pd.DataFrame:
    start_date = date(calendar_year, calendar_month, 1)
    if calendar_month == 12:
        end_date = date(calendar_year + 1, 1, 1)
    else:
        end_date = date(calendar_year, calendar_month + 1, 1)

    sql = """
    SELECT
        po.PONO,
        po.ISSUEDATETIME,
        po.SUPPLIERCODE,
        dbo.GetSSBName(ap.THAINAME) AS VendorName,
        po.PURCHASETYPECODE,
        dbo.GetSSBName(pt.THAINAME) AS PurchaseTypeName,
        po.PURCHASEBUDGETCODE,
        dbo.GetSSBName(bg.THAINAME) AS BudgetName,
        dt.STOCKCODE,
        dbo.GetSSBName(sm.THAINAME) AS ItemName,
        dt.UNITCODE,
        dt.REQUESTQTY,
        dt.LOTPRICE,
        dt.AMT,
        dt.GOODSAMTAFTERITEMDISCOUNT,
        dt.ALLOCATEDVATAMT
    FROM SKPO po
    JOIN SKPODTL dt ON dt.PONO = po.PONO
    LEFT JOIN APMASTER ap ON ap.APCODE = po.SUPPLIERCODE
    LEFT JOIN SYSCONFIG pt ON pt.CODE = po.PURCHASETYPECODE AND pt.CTRLCODE = 100010
    LEFT JOIN SYSCONFIG bg ON bg.CODE = po.PURCHASEBUDGETCODE AND bg.CTRLCODE = 120010
    LEFT JOIN STOCK_MASTER sm ON sm.STOCKCODE = dt.STOCKCODE
    WHERE po.ISSUEDATETIME >= ? AND po.ISSUEDATETIME < ?
      AND po.CXLDATETIME IS NULL
    ORDER BY po.ISSUEDATETIME, po.PONO, dt.SUFFIX
    """
    df = run_query(sql, params=(start_date, end_date))
    if df.empty:
        return df

    df["PODateThai"] = df["ISSUEDATETIME"].apply(to_thai_date)
    df["VendorLabel"] = df["VendorName"].fillna(df["SUPPLIERCODE"])
    df["PurchaseTypeLabel"] = df["PurchaseTypeName"].fillna("(ไม่ระบุวิธี)")
    df["BudgetLabel"] = df["BudgetName"].fillna("(ไม่ระบุแหล่งเงิน)")
    df["ItemLabel"] = df["ItemName"].fillna(df["STOCKCODE"])
    df["NetAmt"] = df["GOODSAMTAFTERITEMDISCOUNT"].fillna(0.0) + df["ALLOCATEDVATAMT"].fillna(0.0)
    return df


_DETAIL_COLUMNS = [
    "ที่", "รหัสวัสดุ", "ชื่อวัสดุ", "หน่วย", "หมวดเงิน", "จำนวน", "ราคา",
    "จำนวนเงิน", "มูลค่าสินค้า", "ภาษีมูลค่าเพิ่ม", "เงินสุทธิ",
]
_WIDTHS = [5, 12, 30, 8, 16, 8, 12, 12, 12, 12, 12]


def build_excel_bytes(df: pd.DataFrame, calendar_year: int, calendar_month: int) -> bytes:
    month_label = THAI_MONTHS[calendar_month]
    fiscal_be = calendar_year + 543
    n_cols = len(_DETAIL_COLUMNS)

    wb = Workbook()
    ws = wb.active
    ws.title = "รายงานซื้อประจำเดือน"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    for i, w in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    c = ws.cell(1, 1, "รายงานสรุปผลการจัดซื้อจัดจ้างในรอบเดือน")
    c.font = Font(name=THAI_FONT, size=16, bold=True)
    c.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    c2 = ws.cell(2, 1, f"ประจำเดือน{month_label} {fiscal_be}")
    c2.font = Font(name=THAI_FONT, size=13)
    c2.alignment = Alignment(horizontal="center")

    row = 4
    if df.empty:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
        ws.cell(row, 1, "ไม่พบข้อมูลตามเงื่อนไขที่เลือก")
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    po_index = 0
    for (pono, issue_date, vendor, ptype), pdf in df.groupby(
        ["PONO", "PODateThai", "VendorLabel", "PurchaseTypeLabel"], sort=False
    ):
        po_index += 1
        po_total = float(pdf["AMT"].fillna(0.0).sum())

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
        header_text = (
            f"{po_index}. เลขที่ใบสั่งซื้อ/ใบแจ้งหนี้ {pono}    วันที่ {issue_date}    "
            f"ผู้ขาย {vendor}    วิธีซื้อหรือจ้าง {ptype}    จำนวนเงิน {po_total:,.2f} บาท"
        )
        hc = ws.cell(row, 1, header_text)
        hc.font = Font(name=THAI_FONT, size=13, bold=True)
        hc.fill = _GROUP_FILL
        row += 1

        for col_idx, label in enumerate(_DETAIL_COLUMNS, start=1):
            cell = ws.cell(row, col_idx, label)
            cell.font = Font(name=THAI_FONT, size=11, bold=True)
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = _BORDER
        row += 1

        for i, (_, r) in enumerate(pdf.iterrows(), start=1):
            values = [
                i, r["STOCKCODE"], r["ItemLabel"], r["UNITCODE"], r["BudgetLabel"],
                r["REQUESTQTY"], r["LOTPRICE"], r["AMT"],
                r["GOODSAMTAFTERITEMDISCOUNT"], r["ALLOCATEDVATAMT"], r["NetAmt"],
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row, col_idx, value)
                cell.font = Font(name=THAI_FONT, size=11)
                cell.border = _BORDER
                if col_idx >= 6:
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
                elif col_idx == 1:
                    cell.alignment = Alignment(horizontal="center")
            row += 1

        row += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
