"""
ตรรกะรายงานค่าเสื่อมราคา (Depreciation) - ใช้ร่วมกันทั้ง 2 รายงาน:
- รายละเอียด ตามแหล่งเงิน (AS_MAS_รายงานค่าเสื่อมราคารายละเอียด_ตามแหล่ง.rpt)
- สรุป ตามกลุ่มงาน (AS_MAS_รายงานค่าเสื่อมราคา_สรุปตามกลุ่มงาน_ปร.rpt)

สูตรคำนวณ ตรวจสอบกับตัวอย่างไฟล์ .xls จริงแล้ว ตรงกันทุกตัวเลข:
- ปีงบประมาณไทย (ต.ค.-ก.ย.) โดย ASMSTYEAR.YEAR = ปี พ.ศ. ของงบ - 543 (เช่น งบ 2569 -> YEAR=2026)
- คอลัมน์ DEPREAMT1..DEPREAMT12 คือค่าเสื่อมรายเดือนแบบ "เดือนงบประมาณ" ไม่ใช่เดือนปฏิทิน
  (ต.ค.=1, พ.ย.=2, ธ.ค.=3, ม.ค.=4, ก.พ.=5, มี.ค.=6, เม.ย.=7, พ.ค.=8, มิ.ย.=9, ก.ค.=10, ส.ค.=11, ก.ย.=12)
- ค่าเสื่อมประจำเดือน = DEPREAMT{n}
- ค่าเสื่อมประจำปี (สะสมในปีงบนั้นจนถึงเดือนที่เลือก) = SUM(DEPREAMT1..DEPREAMT{n})
- ค่าเสื่อมสะสม (ตั้งแต่ได้มา) = BFWVALUEDEPRE (ยกมาต้นปีงบ) + ค่าเสื่อมประจำปี
- มูลค่ารวม = QTY * PRICE, มูลค่าสุทธิ = มูลค่ารวม - ค่าเสื่อมสะสม

ยืนยันด้วยตัวอย่างจริง: ASSETCODE 1100-001-0001/016, งบประมาณ 2569 (=YEAR 2026), ณ สิ้นเดือนสิงหาคม (fiscal month 11)
-> เดือน 71,321.47 / ปี 770,732.02 / สะสม ~8,080,032 / กลุ่ม "เงินงบประมาณ" > "อาคารถาวร"  ตรงกับไฟล์ตัวอย่างทุกค่า
"""
import calendar
import io
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from db_connection import run_query

THAI_MONTHS = [
    "", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
    "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม",
]

THAI_FONT = "TH Sarabun New"
_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
_GROUP_FILL = PatternFill("solid", fgColor="F2F2F2")
_SUBGROUP_FILL = PatternFill("solid", fgColor="FAFAFA")


def to_thai_date(value) -> str:
    if value is None or pd.isna(value):
        return ""
    ts = pd.Timestamp(value)
    return f"{ts.day} {THAI_MONTHS[ts.month]} {ts.year + 543}"


def fiscal_year_and_column(calendar_year: int, calendar_month: int):
    """แปลง (ปี ค.ศ., เดือนปฏิทิน) -> (ASMSTYEAR.YEAR, ลำดับเดือนงบประมาณ 1-12)"""
    if calendar_month >= 10:
        asmst_year = calendar_year + 1
        fiscal_col = calendar_month - 9
    else:
        asmst_year = calendar_year
        fiscal_col = calendar_month + 3
    return asmst_year, fiscal_col


def fiscal_be_label(asmst_year: int) -> str:
    return str(asmst_year + 543)


def fetch_depreciation_data(
    calendar_year: int,
    calendar_month: int,
    division=None,
    dept=None,
    section=None,
) -> pd.DataFrame:
    asmst_year, fiscal_col = fiscal_year_and_column(calendar_year, calendar_month)
    depre_cols = ", ".join(f"y.DEPREAMT{i}" for i in range(1, 13))

    # ตัดครุภัณฑ์ที่ตัดจำหน่ายไปแล้ว ณ วันที่รายงาน (สิ้นเดือนที่เลือก) ออก
    # ถ้าตัดจำหน่ายหลังจากเดือนที่เลือก ยังถือว่าถืออยู่ ณ วันนั้น จึงยังต้องแสดง
    last_day = calendar.monthrange(calendar_year, calendar_month)[1]
    as_of_date = date(calendar_year, calendar_month, last_day)

    conditions = [
        "(cm.DISPOSCODE IS NULL OR (cm.DISPOSDATETIME IS NOT NULL AND cm.DISPOSDATETIME > ?))",
    ]
    params = [as_of_date]

    if division:
        conditions.append("cm.LOCATEDIVISION = ?")
        params.append(division)
    if dept:
        conditions.append("cm.LOCATEDEPT = ?")
        params.append(dept)
    if section:
        conditions.append("cm.LOCATESECTION = ?")
        params.append(section)

    where_clause = " AND ".join(conditions) if conditions else "1=1"
    depre_sum_cols = ", ".join(f"SUM(DEPREAMT{i}) AS DEPREAMT{i}" for i in range(1, 13))

    # ASMSTYEAR เก็บแยกได้หลายแถวต่อ ASSETCODE+SUFFIX+YEAR ถ้าค่าเสื่อมถูกปันส่วนข้าม GLDEPT
    # (พบ 13 รายการจากทั้งหมดหมื่นกว่ารายการ) ต้องรวมยอดก่อน join ไม่งั้นจะได้แถวซ้ำ
    sql = f"""
    WITH YearAgg AS (
        SELECT ASSETCODE, SUFFIX, SUM(BFWVALUEDEPRE) AS BFWVALUEDEPRE, {depre_sum_cols}
        FROM ASMSTYEAR
        WHERE YEAR = ?
        GROUP BY ASSETCODE, SUFFIX
    )
    SELECT
        cm.ASSETCODE,
        cm.SUFFIX,
        dbo.GetSSBName(ISNULL(cm.THAINAME, m.THAINAME)) AS AssetName,
        COALESCE(cm.ACQDATETIME, m.FIRSTDATETIME) AS AcqDate,
        cm.QTY,
        cm.PRICE,
        cm.LOCATEDIVISION,
        div.DivisionName,
        cm.ARTICLEGROUP,
        dbo.GetSSBName(ag.THAINAME) AS ArticleGroupName,
        cm.PURCHASEBUDGETCODE,
        dbo.GetSSBName(bg.THAINAME) AS BudgetName,
        y.BFWVALUEDEPRE,
        {depre_cols}
    FROM YearAgg y
    JOIN ASMSTCM cm ON cm.ASSETCODE = y.ASSETCODE AND cm.SUFFIX = y.SUFFIX
    JOIN ASMST m ON m.ASSETCODE = cm.ASSETCODE
    LEFT JOIN Division div ON cm.LOCATEDIVISION = div.Division
    LEFT JOIN SYSCONFIG ag ON ag.CODE = cm.ARTICLEGROUP AND ag.CTRLCODE = 100012
    LEFT JOIN SYSCONFIG bg ON bg.CODE = cm.PURCHASEBUDGETCODE AND bg.CTRLCODE = 120010
    WHERE {where_clause}
    ORDER BY cm.PURCHASEBUDGETCODE, cm.ARTICLEGROUP, cm.ASSETCODE
    """
    params = [asmst_year] + params

    df = run_query(sql, params=tuple(params))
    if df.empty:
        return df

    depre_amt_cols = [f"DEPREAMT{i}" for i in range(1, fiscal_col + 1)]
    df["DepreMonth"] = df[f"DEPREAMT{fiscal_col}"].fillna(0.0)
    df["DepreYearCum"] = df[depre_amt_cols].fillna(0.0).sum(axis=1)
    df["AccumDepre"] = df["BFWVALUEDEPRE"].fillna(0.0) + df["DepreYearCum"]
    df["TotalValue"] = df["QTY"].fillna(0.0) * df["PRICE"].fillna(0.0)
    df["NetValue"] = df["TotalValue"] - df["AccumDepre"]
    df["AcqDateThai"] = df["AcqDate"].apply(to_thai_date)
    df["BudgetLabel"] = df["BudgetName"].fillna("(ไม่ระบุแหล่งเงิน)")
    df["ArticleGroupLabel"] = (
        "(" + df["ARTICLEGROUP"].fillna("") + ") " + df["ArticleGroupName"].fillna("")
    )
    df["DivisionLabel"] = df["DivisionName"].fillna("(ไม่ระบุฝ่าย)")
    return df


_DETAIL_COLUMNS = ["รหัส", "รายการ", "วันเดือนปี", "จำนวน", "มูลค่ารวม",
                    "ค่าเสื่อม ประจำเดือน", "ค่าเสื่อม ประจำปี", "สะสม", "มูลค่าสุทธิ"]
_DETAIL_WIDTHS = [16, 38, 14, 8, 14, 14, 14, 14, 14]


def _sheet_header(ws, n_cols, title, fiscal_be, month_label, print_dt):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    c = ws.cell(1, 1, title)
    c.font = Font(name=THAI_FONT, size=18, bold=True)
    c.alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    c2 = ws.cell(2, 1, f"รายงานทรัพย์สิน ณ สิ้นเดือน{month_label} {fiscal_be}")
    c2.font = Font(name=THAI_FONT, size=13)
    c2.alignment = Alignment(horizontal="center")
    return 4


def build_detail_excel_bytes(df: pd.DataFrame, calendar_year: int, calendar_month: int) -> bytes:
    asmst_year, _ = fiscal_year_and_column(calendar_year, calendar_month)
    fiscal_be = fiscal_be_label(asmst_year)
    month_label = THAI_MONTHS[calendar_month]

    wb = Workbook()
    ws = wb.active
    ws.title = "รายละเอียดค่าเสื่อม"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    n_cols = len(_DETAIL_COLUMNS)
    for i, w in enumerate(_DETAIL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row = _sheet_header(ws, n_cols, "รายงานค่าเสื่อมราคา รายละเอียด ตามแหล่งเงิน", fiscal_be, month_label, None)

    if df.empty:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
        ws.cell(row, 1, "ไม่พบข้อมูลตามเงื่อนไขที่เลือก")
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    grand_total_value = 0.0
    grand_total_depre = 0.0
    grand_total_net = 0.0
    grand_count = 0

    for budget_label, bdf in df.groupby("BudgetLabel", sort=False):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
        gc = ws.cell(row, 1, budget_label)
        gc.font = Font(name=THAI_FONT, size=14, bold=True)
        gc.fill = _GROUP_FILL
        row += 1

        for art_label, adf in bdf.groupby("ArticleGroupLabel", sort=False):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
            ac = ws.cell(row, 1, art_label)
            ac.font = Font(name=THAI_FONT, size=13, bold=True)
            ac.fill = _SUBGROUP_FILL
            row += 1

            for col_idx, label in enumerate(_DETAIL_COLUMNS, start=1):
                cell = ws.cell(row, col_idx, label)
                cell.font = Font(name=THAI_FONT, size=12, bold=True)
                cell.fill = _HEADER_FILL
                cell.alignment = Alignment(horizontal="center")
                cell.border = _BORDER
            row += 1

            sub_total_value = 0.0
            sub_total_depre = 0.0
            sub_total_net = 0.0
            for _, r in adf.iterrows():
                values = [
                    r["ASSETCODE"],
                    r["AssetName"],
                    r["AcqDateThai"],
                    r["QTY"],
                    r["TotalValue"],
                    r["DepreMonth"],
                    r["DepreYearCum"],
                    r["AccumDepre"],
                    r["NetValue"],
                ]
                for col_idx, value in enumerate(values, start=1):
                    cell = ws.cell(row, col_idx, value)
                    cell.font = Font(name=THAI_FONT, size=12)
                    cell.border = _BORDER
                    if col_idx >= 5:
                        cell.number_format = "#,##0.00"
                        cell.alignment = Alignment(horizontal="right")
                    elif col_idx == 4:
                        cell.alignment = Alignment(horizontal="center")
                sub_total_value += float(r["TotalValue"])
                sub_total_depre += float(r["AccumDepre"])
                sub_total_net += float(r["NetValue"])
                row += 1

            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            tc = ws.cell(row, 1, f"รวม {art_label} ({len(adf)} รายการ)")
            tc.font = Font(name=THAI_FONT, size=12, bold=True)
            tc.alignment = Alignment(horizontal="right")
            for col_idx, val in [(5, sub_total_value), (8, sub_total_depre), (9, sub_total_net)]:
                c = ws.cell(row, col_idx, val)
                c.font = Font(name=THAI_FONT, size=12, bold=True)
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")
            row += 2

            grand_total_value += sub_total_value
            grand_total_depre += sub_total_depre
            grand_total_net += sub_total_net
            grand_count += len(adf)

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    gcell = ws.cell(row, 1, f"รวมทั้งสิ้น ({grand_count} รายการ)")
    gcell.font = Font(name=THAI_FONT, size=14, bold=True)
    gcell.alignment = Alignment(horizontal="right")
    for col_idx, val in [(5, grand_total_value), (8, grand_total_depre), (9, grand_total_net)]:
        c = ws.cell(row, col_idx, val)
        c.font = Font(name=THAI_FONT, size=14, bold=True)
        c.number_format = "#,##0.00"
        c.alignment = Alignment(horizontal="right")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


_SUMMARY_COLUMNS = ["ฝ่าย/หน่วยงาน", "หมวดครุภัณฑ์", "จำนวนรายการ", "มูลค่ารวม",
                     "ค่าเสื่อม ประจำเดือน", "ค่าเสื่อม ประจำปี", "สะสม", "มูลค่าสุทธิ"]
_SUMMARY_WIDTHS = [30, 30, 12, 16, 16, 16, 16, 16]


def build_summary_excel_bytes(df: pd.DataFrame, calendar_year: int, calendar_month: int) -> bytes:
    asmst_year, _ = fiscal_year_and_column(calendar_year, calendar_month)
    fiscal_be = fiscal_be_label(asmst_year)
    month_label = THAI_MONTHS[calendar_month]

    wb = Workbook()
    ws = wb.active
    ws.title = "สรุปค่าเสื่อมตามกลุ่มงาน"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    n_cols = len(_SUMMARY_COLUMNS)
    for i, w in enumerate(_SUMMARY_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row = _sheet_header(ws, n_cols, "รายงานค่าเสื่อมราคา สรุปตามกลุ่มงาน", fiscal_be, month_label, None)

    if df.empty:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
        ws.cell(row, 1, "ไม่พบข้อมูลตามเงื่อนไขที่เลือก")
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    for col_idx, label in enumerate(_SUMMARY_COLUMNS, start=1):
        cell = ws.cell(row, col_idx, label)
        cell.font = Font(name=THAI_FONT, size=13, bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = _BORDER
    row += 1

    grand = dict(count=0, value=0.0, month=0.0, year=0.0, accum=0.0, net=0.0)

    grouped = df.groupby(["DivisionLabel", "ArticleGroupLabel"], sort=False)
    agg = grouped.agg(
        count=("ASSETCODE", "count"),
        TotalValue=("TotalValue", "sum"),
        DepreMonth=("DepreMonth", "sum"),
        DepreYearCum=("DepreYearCum", "sum"),
        AccumDepre=("AccumDepre", "sum"),
        NetValue=("NetValue", "sum"),
    ).reset_index()

    for _, r in agg.iterrows():
        values = [
            r["DivisionLabel"], r["ArticleGroupLabel"], r["count"],
            r["TotalValue"], r["DepreMonth"], r["DepreYearCum"], r["AccumDepre"], r["NetValue"],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row, col_idx, value)
            cell.font = Font(name=THAI_FONT, size=13)
            cell.border = _BORDER
            if col_idx >= 4:
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            elif col_idx == 3:
                cell.alignment = Alignment(horizontal="center")
        row += 1
        grand["count"] += r["count"]
        grand["value"] += r["TotalValue"]
        grand["month"] += r["DepreMonth"]
        grand["year"] += r["DepreYearCum"]
        grand["accum"] += r["AccumDepre"]
        grand["net"] += r["NetValue"]

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    gcell = ws.cell(row, 1, "รวมทั้งสิ้น")
    gcell.font = Font(name=THAI_FONT, size=14, bold=True)
    gcell.alignment = Alignment(horizontal="right")
    for col_idx, val in [(3, grand["count"]), (4, grand["value"]), (5, grand["month"]),
                          (6, grand["year"]), (7, grand["accum"]), (8, grand["net"])]:
        c = ws.cell(row, col_idx, val)
        c.font = Font(name=THAI_FONT, size=14, bold=True)
        if col_idx != 3:
            c.number_format = "#,##0.00"
        c.alignment = Alignment(horizontal="right")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
